import sys
import types

# Stub only the optional SDK at import time: these tests replace the model calls.
google_pkg = types.ModuleType("google")
google_pkg.__path__ = []
genai_pkg = types.ModuleType("google.genai")
genai_pkg.__path__ = []
errors_mod = types.ModuleType("google.genai.errors")
class APIError(Exception):
    pass
errors_mod.APIError = APIError
types_mod = types.ModuleType("google.genai.types")
types_mod.Part = type("Part", (), {})
types_mod.GenerateContentConfig = type("GenerateContentConfig", (), {})
genai_pkg.errors = errors_mod
genai_pkg.types = types_mod
google_pkg.genai = genai_pkg
sys.modules.setdefault("google", google_pkg)
sys.modules.setdefault("google.genai", genai_pkg)
sys.modules.setdefault("google.genai.errors", errors_mod)
sys.modules.setdefault("google.genai.types", types_mod)

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from titivillus.api import app
from titivillus import pipeline
from titivillus.config import Settings, MODO_CAPA, MODO_FICHA, MODO_NUNCA
from titivillus.planilha import exportar_excel

client = TestClient(app)

def test_api_health_and_settings_hide_secrets():
    assert client.get('/api/health').json()['ok'] is True
    data = client.get('/api/settings').json()
    assert data['covers_dir'].endswith('/capas')
    assert 'GEMINI_API_KEY' not in str(data)

def test_file_listing_modes():
    covers = client.get('/api/files?mode=capa')
    cards = client.get('/api/files?mode=ficha')
    assert covers.status_code == cards.status_code == 200
    assert covers.json()['count'] >= 0 and cards.json()['count'] >= 0
    assert covers.json()['mode'] == 'capa' and cards.json()['mode'] == 'ficha'

def test_run_requires_gemini_key_when_missing():
    response = client.post('/api/run', json={'mode':'capa'})
    # CI environment should not depend on a personal .env key.
    if not __import__('titivillus.api', fromlist=['API_KEY']).API_KEY:
        assert response.status_code == 400
        assert 'GEMINI_API_KEY' in response.json()['detail']

def test_excel_export_does_not_touch_project_files(tmp_path):
    output = tmp_path / 'test-catalog.xlsx'
    exportar_excel([{'titulo':'Livro Teste','volume':'1','autor':'Autora','faixa_etaria':'Livre',
        'titulo_google_books':'Livro Teste','autor_google_books':'Autora','validado':True,
        'status':'validado','fonte_validacao':'Google Books','arquivo':'capa.jpg','erro':None}], output, MODO_CAPA)
    wb = load_workbook(output, read_only=True)
    assert wb.active.max_row == 2
    assert wb.active['A2'].value == 'Livro Teste'
    wb.close()

def test_pipeline_processes_cover_and_card_jobs_and_exports(tmp_path, monkeypatch):
    covers = tmp_path / 'covers'; covers.mkdir(); (covers / 'sample.jpg').write_bytes(b'fake image')
    cards = tmp_path / 'cards'; cards.mkdir(); (cards / 'card.jpg').write_bytes(b'fake image')
    cover_out = tmp_path / 'covers.xlsx'; card_out = tmp_path / 'cards.xlsx'
    monkeypatch.setattr(pipeline, 'processar_imagem', lambda path, policy: {
        'titulo':'Obra de Teste','volume':'1','autor':'Pessoa','faixa_etaria':'Livre',
        'titulo_google_books':'Obra de Teste','autor_google_books':'Pessoa','validado':True,
        'status':'validado','fonte_validacao':'Google Books','arquivo':path.name,'erro':None})
    monkeypatch.setattr(pipeline, 'processar_ficha', lambda path: {
        'arquivo':path.name,'titulo':'Ficha de Teste','autores':'Pessoa','isbn':'123',
        'status':'lida','erro':None})
    cover_settings = Settings(modo=MODO_CAPA, pasta_capas=covers, arquivo_excel=cover_out,
        pasta_fichas=cards, arquivo_excel_fichas=card_out, modo_renomear=MODO_NUNCA)
    cover_records = pipeline.executar(cover_settings)
    assert len(cover_records) == 1 and cover_out.exists()
    card_settings = Settings(modo=MODO_FICHA, pasta_capas=covers, arquivo_excel=cover_out,
        pasta_fichas=cards, arquivo_excel_fichas=card_out, modo_renomear=MODO_NUNCA)
    card_records = pipeline.executar(card_settings)
    assert len(card_records) == 1 and card_out.exists()
    for path in (cover_out, card_out):
        wb = load_workbook(path, read_only=True)
        assert wb.active.max_row == 2
        wb.close()

def test_real_cover_processing_function_with_mocked_external_services(tmp_path, monkeypatch):
    image = tmp_path / 'cover.jpg'; image.write_bytes(b'fake image')
    monkeypatch.setattr(pipeline, 'ler_capa', lambda path: {'titulo':'Livro Mock','volume':'2','autor':'Autora'})
    monkeypatch.setattr(pipeline, 'validar', lambda data: {**data,'validado':True,'status':'validado','titulo_google_books':'Livro Mock','autor_google_books':'Autora','fonte_validacao':'Google Books'})
    monkeypatch.setattr(pipeline, 'ler_faixa_etaria', lambda title, author: '12+')
    result = pipeline.processar_imagem(image, MODO_NUNCA)
    assert result['titulo'] == 'Livro Mock'
    assert result['faixa_etaria'] == '12+' and result['validado'] is True

def test_real_card_processing_function_with_mocked_gemini(tmp_path, monkeypatch):
    image = tmp_path / 'card.jpg'; image.write_bytes(b'fake image')
    monkeypatch.setattr(pipeline, 'ler_ficha', lambda path: {'titulo':'Livro da Ficha','isbn':'9780000000000','autores':['Autora']})
    result = pipeline.processar_ficha(image)
    assert result['titulo'] == 'Livro da Ficha'
    assert result['status'] == 'lida'
