"""Ordenação alfabética e exportação para Excel (planilha formatada)."""

import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .config import MODO_CAPA, MODO_FICHA
from .textos import normalizar

COR_CABECALHO = "1F3864"
COR_ZEBRA = "F2F5FA"
COR_LINHA = "D0D0D0"
FONTE = "Arial"

# Palavras que ficam em minúscula ao converter TÍTULO EM CAIXA ALTA.
_MINUSCULAS = {"a", "o", "as", "os", "de", "da", "do", "das", "dos", "e", "em", "na",
               "no", "nas", "nos", "para", "por", "com", "um", "uma", "ao", "à"}


# --------------------------------------------------------------------------- #
# Esquemas de coluna
# --------------------------------------------------------------------------- #
# (campo, cabeçalho, largura, alinhamento horizontal, quebra de linha, formato)
# Os campos de `colunas` seguem a ordem exata da planilha. `extras` só entram
# (no fim) se algum registro tiver conteúdo neles, para nunca perder dado.

ESQUEMAS = {
    MODO_CAPA: {
        "aba": "Livros",
        "altura_linha": 30,
        "colunas": [
            ("titulo", "Título", 44, "left", True, None),
            ("volume", "Vol.", 6, "center", False, None),
            ("autor", "Autor", 28, "left", False, None),
            ("faixa_etaria", "Faixa etária", 12, "center", False, None),
            ("status", "Status", 34, "left", False, None),
            ("validado", "Validado", 10, "center", False, None),
            ("titulo_google_books", "Título (Google Books)", 30, "left", False, None),
            ("autor_google_books", "Autor (Google Books)", 26, "left", False, None),
            ("fonte_validacao", "Fonte da validação", 20, "left", False, None),
            ("arquivo", "Arquivo", 60, "left", False, None),
            ("erro", "Erro", 70, "left", False, None),
        ],
        "extras": [],
    },
    MODO_FICHA: {
        "aba": "Fichas",
        "altura_linha": 45,
        "colunas": [
            ("titulo", "Título", 30, "left", True, None),
            ("autores", "Autores", 32, "left", True, None),
            ("tradutor", "Tradutor", 18, "left", False, None),
            ("ilustrador", "Ilustrador", 18, "left", False, None),
            ("editora", "Editora", 18, "left", False, None),
            ("cidade", "Cidade", 15, "left", False, None),
            ("ano_publicacao", "Ano", 8, "center", False, None),
            ("edicao", "Edição", 9, "center", False, None),
            ("serie_colecao", "Série / Coleção", 26, "left", True, None),
            ("volume", "Vol.", 6, "center", False, None),
            ("titulo_original", "Título original", 32, "left", True, None),
            ("isbn", "ISBN", 36, "left", True, "@"),
            ("paginas", "Págs.", 7, "center", False, None),
            ("idioma", "Idioma", 11, "center", False, None),
            ("cdd", "CDD", 8, "center", False, "@"),
            ("cdu", "CDU", 8, "center", False, "@"),
            ("assuntos", "Assuntos", 48, "left", True, None),
            ("bibliotecario_crb", "Bibliotecário (CRB)", 28, "left", False, None),
            ("status", "Status", 9, "center", False, None),
            ("outras_informacoes", "Outras informações", 20, "left", False, None),
            ("arquivo", "Arquivo", 40, "left", False, None),
            ("ficha_catalografica", "Ficha catalográfica", 60, "left", False, None),
        ],
        "extras": [
            ("subtitulo", "Subtítulo", 30, "left", True, None),
            ("organizador", "Organizador", 20, "left", False, None),
            ("erro", "Erro", 50, "left", False, None),
        ],
    },
}


# --------------------------------------------------------------------------- #
# Tratamento dos valores
# --------------------------------------------------------------------------- #
def _vazio(valor):
    if valor is None:
        return True
    if isinstance(valor, float) and valor != valor:  # NaN
        return True
    return isinstance(valor, str) and not valor.strip()


def _numero_volume(valor):
    if isinstance(valor, str):
        achado = re.search(r"\d+", valor)
        if achado:
            return int(achado.group())
    elif isinstance(valor, int) and not isinstance(valor, bool):
        return valor
    return 0


def _caixa_alta_para_titulo(texto):
    """'ONE PIECE' -> 'One Piece'. Só mexe em texto 100% maiúsculo; o resto
    (que já tem caixa mista) fica exatamente como veio. Siglas curtas sem
    vogal (TCP, HTML) são preservadas."""
    letras = [c for c in texto if c.isalpha()]
    if not letras or not all(c.isupper() for c in letras):
        return texto
    palavras = texto.split(" ")
    saida = []
    for i, palavra in enumerate(palavras):
        puro = re.sub(r"\W", "", palavra)
        if puro and len(puro) <= 4 and not re.search(r"[AEIOUÁÀÂÃÉÊÍÓÔÕÚ]", puro):
            saida.append(palavra)
        elif i > 0 and palavra.lower() in _MINUSCULAS:
            saida.append(palavra.lower())
        else:
            saida.append(palavra.capitalize())
    return " ".join(saida)


def _valor_celula(campo, valor, modo):
    if _vazio(valor):
        return None
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, list):
        valor = "; ".join(str(v).strip() for v in valor if not _vazio(v))
        if not valor:
            return None
    if isinstance(valor, str):
        valor = ILLEGAL_CHARACTERS_RE.sub("", valor).strip()
        if modo == MODO_CAPA:
            if campo in ("titulo", "autor"):
                valor = _caixa_alta_para_titulo(valor)
            elif campo == "volume" and valor.isdigit():
                return int(valor)  # volume numérico vira número, como na planilha
    return valor


def _colunas(modo, registros):
    esquema = ESQUEMAS[modo]
    colunas = list(esquema["colunas"])
    for extra in esquema["extras"]:
        if any(not _vazio(r.get(extra[0])) for r in registros):
            colunas.append(extra)
    return colunas


def _ordenar(registros):
    def chave(r):
        titulo = r.get("titulo")
        sem_titulo = not (isinstance(titulo, str) and titulo.strip())
        return (sem_titulo, normalizar(titulo), _numero_volume(r.get("volume")))

    return sorted(registros, key=chave)  # sorted() é estável


# --------------------------------------------------------------------------- #
# Exportação
# --------------------------------------------------------------------------- #
def exportar_excel(registros, caminho, modo=MODO_CAPA):
    esquema = ESQUEMAS[modo]
    colunas = _colunas(modo, registros)
    registros = _ordenar(registros)

    wb = Workbook()
    ws = wb.active
    ws.title = esquema["aba"]

    fonte = Font(name=FONTE, size=10)
    fonte_negrito = Font(name=FONTE, size=10, bold=True)
    borda = Border(bottom=Side(style="thin", color=COR_LINHA))
    zebra = PatternFill("solid", fgColor=COR_ZEBRA)

    # Cabeçalho
    cab_fonte = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor=COR_CABECALHO)
    cab_alin = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j, (_, titulo, largura, *_resto) in enumerate(colunas, start=1):
        celula = ws.cell(row=1, column=j, value=titulo)
        celula.font, celula.fill, celula.alignment = cab_fonte, cab_fill, cab_alin
        ws.column_dimensions[celula.column_letter].width = largura
    ws.row_dimensions[1].height = 30

    # Dados
    for i, registro in enumerate(registros, start=2):
        ws.row_dimensions[i].height = esquema["altura_linha"]
        for j, (campo, _, _, horizontal, quebra, formato) in enumerate(
            colunas, start=1
        ):
            valor = _valor_celula(campo, registro.get(campo), modo)
            if formato == "@" and valor is not None:
                valor = str(valor)
            celula = ws.cell(row=i, column=j, value=valor)
            celula.font = fonte_negrito if campo == "titulo" else fonte
            celula.alignment = Alignment(
                horizontal=horizontal, vertical="top", wrap_text=quebra or None
            )
            celula.border = borda
            if formato:
                celula.number_format = formato
            if i % 2 == 1:  # linhas 3, 5, 7...
                celula.fill = zebra

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    try:
        wb.save(caminho)
    except PermissionError as erro:
        raise RuntimeError(
            f"Sem permissão para gravar {caminho}. "
            "Feche a planilha se ela estiver aberta."
        ) from erro
